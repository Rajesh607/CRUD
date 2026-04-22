import json
from pathlib import Path

import openpyxl
from docx import Document
from PyPDF2 import PdfReader

ALLOWED_EXTENSIONS = {".doc", ".docx", ".pdf", ".xlsx", ".xls", ".json", ".txt"}


def detect_file_type(filename: str) -> str:
    return Path(filename).suffix.lower()


def validate_file_type(file_ext: str) -> bool:
    return file_ext in ALLOWED_EXTENSIONS


def extract_content(file_path: str, file_ext: str) -> tuple[str, dict]:
    if file_ext == ".pdf":
        reader = PdfReader(file_path)
        text = "\n".join([page.extract_text() or "" for page in reader.pages])
        return text, {"pages": len(reader.pages)}

    if file_ext in {".doc", ".docx"}:
        doc = Document(file_path)
        text = "\n".join([p.text for p in doc.paragraphs])
        return text, {"paragraphs": len(doc.paragraphs)}

    if file_ext in {".xlsx", ".xls"}:
        wb = openpyxl.load_workbook(file_path)
        values = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values.append([v for v in row if v is not None])
        text = "\n".join([", ".join(map(str, row)) for row in values if row])
        return text, {"sheets": len(wb.sheetnames)}

    if file_ext == ".json":
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        text = json.dumps(data)
        return text, {"json_keys": list(data.keys()) if isinstance(data, dict) else []}

    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    return text, {"length": len(text)}
