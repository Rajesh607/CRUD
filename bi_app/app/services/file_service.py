import json
import os
from pathlib import Path

import openpyxl
from docx import Document
from PyPDF2 import PdfReader

ALLOWED_EXTENSIONS = {".doc", ".docx", ".pdf", ".xlsx", ".xls", ".json", ".txt"}


def detect_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def validate_file(extension: str, size_bytes: int, max_mb: int) -> None:
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported extension: {extension}")
    if size_bytes > max_mb * 1024 * 1024:
        raise ValueError(f"File exceeds {max_mb} MB")


def extract_file(path: str, extension: str) -> tuple[str, dict]:
    if extension == ".pdf":
        reader = PdfReader(path)
        text = "\n".join([p.extract_text() or "" for p in reader.pages])
        return text, {"pages": len(reader.pages)}
    if extension in {".doc", ".docx"}:
        doc = Document(path)
        text = "\n".join([p.text for p in doc.paragraphs])
        return text, {"paragraphs": len(doc.paragraphs)}
    if extension in {".xlsx", ".xls"}:
        wb = openpyxl.load_workbook(path)
        rows: list[str] = []
        for sh in wb.worksheets:
            for row in sh.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None]
                if vals:
                    rows.append(", ".join(vals))
        return "\n".join(rows), {"sheets": len(wb.worksheets)}
    if extension == ".json":
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return json.dumps(data), {"kind": type(data).__name__}

    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    return text, {"bytes": os.path.getsize(path)}
